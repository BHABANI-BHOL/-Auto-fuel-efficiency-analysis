"""
Automotive Fuel Efficiency & Performance Analysis
--------------------------------------------------
End-to-end analysis: load -> clean -> engineer features -> analyze -> visualize -> export.

Dataset : "cars" (406 vehicles, model years 1970-1982), a public dataset shipped with the
          `vega_datasets` package (originally from the StatLib / UCI Auto-MPG data).
Run     : python analysis.py
Outputs : data/cars_clean.csv, charts/*.png, output/Auto_Efficiency_Dashboard.xlsx, output/metrics.json
"""
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA, CHARTS, OUT = ROOT / "data", ROOT / "charts", ROOT / "output"
for p in (DATA, CHARTS, OUT):
    p.mkdir(exist_ok=True)

# ----------------------------------------------------------------------------
# 1. LOAD
# ----------------------------------------------------------------------------
raw_path = DATA / "cars_raw.csv"
if not raw_path.exists():
    from vega_datasets import data as vega

    vega.cars().to_csv(raw_path, index=False)
raw = pd.read_csv(raw_path)
print(f"Loaded raw data: {raw.shape[0]} rows x {raw.shape[1]} columns")

# ----------------------------------------------------------------------------
# 2. CLEAN
# ----------------------------------------------------------------------------
df = raw.copy()
df.columns = ["name", "mpg", "cylinders", "displacement", "horsepower",
              "weight_lbs", "acceleration", "year", "origin"]

quality = {
    "raw_rows": int(len(df)),
    "missing_mpg": int(df["mpg"].isna().sum()),
    "missing_horsepower": int(df["horsepower"].isna().sum()),
    "duplicate_rows": int(df.duplicated().sum()),
}

df["model_year"] = pd.to_datetime(df["year"]).dt.year           # parse date -> model year
df = df.drop(columns="year")
df["name"] = df["name"].str.strip().str.lower()

# Brand extraction + fix inconsistent spellings
MAKE_FIX = {"chevy": "chevrolet", "chevroelt": "chevrolet", "maxda": "mazda",
            "mercedes-benz": "mercedes", "toyouta": "toyota",
            "vokswagen": "volkswagen", "vw": "volkswagen"}
df["make"] = df["name"].str.split().str[0].replace(MAKE_FIX)

# MPG is the target metric -> rows without it cannot be analysed
df = df.dropna(subset=["mpg"]).copy()
# Horsepower: impute the few gaps with the median of the same cylinder group
df["horsepower"] = df["horsepower"].fillna(df.groupby("cylinders")["horsepower"].transform("median"))
quality["clean_rows"] = int(len(df))

# ----------------------------------------------------------------------------
# 3. FEATURE ENGINEERING
# ----------------------------------------------------------------------------
df["power_to_weight"] = (df["horsepower"] / df["weight_lbs"] * 1000).round(2)   # HP per 1,000 lbs
df["efficiency_class"] = pd.cut(df["mpg"], [0, 20, 30, 100],
                                labels=["Low (<20)", "Medium (20-30)", "High (30+)"], right=False)
df["era"] = pd.cut(df["model_year"], [1969, 1973, 1977, 1982],
                   labels=["1970-73", "1974-77", "1978-82"])
df["high_eff"] = (df["mpg"] >= 30).astype(int)
df.to_csv(DATA / "cars_clean.csv", index=False)

# ----------------------------------------------------------------------------
# 4. ANALYZE
# ----------------------------------------------------------------------------
origin_order = ["USA", "Japan", "Europe"]
by_origin = (df.groupby("origin")
             .agg(vehicles=("mpg", "size"), avg_mpg=("mpg", "mean"), avg_hp=("horsepower", "mean"),
                  avg_weight=("weight_lbs", "mean"), high_eff_share=("high_eff", "mean"))
             .reindex(origin_order))

trend = df.pivot_table(index="model_year", columns="origin", values="mpg", aggfunc="mean")[origin_order]
overall_trend = df.groupby("model_year")["mpg"].mean()

by_cyl = df[df["cylinders"].isin([4, 6, 8])].groupby("cylinders")["mpg"].agg(["size", "mean"])
by_era_origin = df.pivot_table(index="era", columns="origin", values="high_eff", aggfunc="mean",
                               observed=True)[origin_order]
top_makes = (df.groupby("make").agg(n=("mpg", "size"), avg_mpg=("mpg", "mean"))
             .query("n >= 5").sort_values("avg_mpg", ascending=False))

corr = df[["mpg", "weight_lbs", "displacement", "horsepower", "acceleration"]].corr()
slope, intercept = np.polyfit(df["weight_lbs"], df["mpg"], 1)

first, last = overall_trend.iloc[0], overall_trend.iloc[-1]
heavy_q = df["weight_lbs"].quantile(0.75)
light_q = df["weight_lbs"].quantile(0.25)
mpg_heavy = df.loc[df["weight_lbs"] >= heavy_q, "mpg"].mean()
mpg_light = df.loc[df["weight_lbs"] <= light_q, "mpg"].mean()

metrics = {
    "data_quality": quality,
    "avg_mpg_overall": round(df["mpg"].mean(), 1),
    "avg_mpg_by_origin": by_origin["avg_mpg"].round(1).to_dict(),
    "high_eff_share_by_origin_pct": (by_origin["high_eff_share"] * 100).round(1).to_dict(),
    "mpg_1970": round(first, 1), "mpg_1982": round(last, 1),
    "mpg_improvement_pct": round((last / first - 1) * 100, 1),
    "mpg_by_cylinders": by_cyl["mean"].round(1).to_dict(),
    "corr_mpg_weight": round(corr.loc["mpg", "weight_lbs"], 2),
    "corr_mpg_displacement": round(corr.loc["mpg", "displacement"], 2),
    "corr_mpg_horsepower": round(corr.loc["mpg", "horsepower"], 2),
    "mpg_drop_per_1000lbs": round(-slope * 1000, 1),
    "r_squared_weight": round(corr.loc["mpg", "weight_lbs"] ** 2, 2),
    "mpg_lightest_quartile": round(mpg_light, 1),
    "mpg_heaviest_quartile": round(mpg_heavy, 1),
    "high_eff_share_by_era_pct": (by_era_origin * 100).round(1).to_dict(),
    "top_make": top_makes.index[0], "top_make_mpg": round(top_makes.iloc[0]["avg_mpg"], 1),
    "bottom_make": top_makes.index[-1], "bottom_make_mpg": round(top_makes.iloc[-1]["avg_mpg"], 1),
}
(OUT / "metrics.json").write_text(json.dumps(metrics, indent=2, default=str))

# ----------------------------------------------------------------------------
# 5. VISUALIZE  (colour-blind-safe palette, thin marks, direct labels)
# ----------------------------------------------------------------------------
for f in ("Inter-Regular.otf", "Inter-SemiBold.otf", "Inter-Bold.otf"):
    fp = Path("/usr/share/fonts/opentype/inter") / f
    if fp.exists():
        font_manager.fontManager.addfont(str(fp))
SURFACE, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e6e5e1"
COL = {"USA": "#2a78d6", "Japan": "#eb6834", "Europe": "#1baf7a"}
plt.rcParams.update({
    "font.family": ["Inter", "DejaVu Sans"], "figure.facecolor": SURFACE, "axes.facecolor": SURFACE,
    "axes.edgecolor": GRID, "axes.labelcolor": INK2, "text.color": INK, "xtick.color": INK2,
    "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8, "axes.axisbelow": True,
    "axes.titlesize": 13, "axes.titleweight": "bold", "axes.titlelocation": "left",
    "axes.titlecolor": INK, "legend.frameon": False, "figure.dpi": 130,
})

def save(fig, name):
    fig.tight_layout()
    fig.savefig(CHARTS / name, facecolor=SURFACE)
    plt.close(fig)

# 5.1 MPG trend by origin
fig, ax = plt.subplots(figsize=(8, 4.6))
for o in origin_order:
    s = trend[o].dropna()
    ax.plot(s.index, s.values, color=COL[o], lw=2, marker="o", ms=4.5, label=o,
            markeredgecolor=SURFACE, markeredgewidth=1.2)
    dy = {"Europe": 6, "Japan": -6, "USA": 0}[o]
    ax.annotate(o, (s.index[-1], s.values[-1]), xytext=(8, dy), textcoords="offset points",
                va="center", color=INK2, fontsize=9, fontweight="bold")
ax.set_title("Average fuel efficiency rose ~75% between 1970 and 1982")
ax.title.set_fontsize(11.5)
ax.set_ylabel("Average miles per gallon")
ax.set_xlim(1969.6, 1983.4)
ax.set_xticks(range(1970, 1983, 2))
ax.legend(loc="upper left", ncols=3)
save(fig, "01_mpg_trend_by_origin.png")

# 5.2 Weight vs MPG scatter
fig, ax = plt.subplots(figsize=(8, 4.8))
for o in origin_order:
    d = df[df["origin"] == o]
    ax.scatter(d["weight_lbs"], d["mpg"], s=22, color=COL[o], alpha=0.75, label=o,
               edgecolor=SURFACE, linewidth=0.6)
xs = np.linspace(df["weight_lbs"].min(), df["weight_lbs"].max(), 50)
ax.plot(xs, slope * xs + intercept, color=INK, lw=1.4, ls="--", label="Linear trend")
ax.set_title(f"Every extra 1,000 lbs costs about {-slope*1000:.1f} MPG (r = {corr.loc['mpg','weight_lbs']:.2f})")
ax.title.set_fontsize(11.5)
ax.set_xlabel("Vehicle weight (lbs)")
ax.set_ylabel("Miles per gallon")
ax.legend(loc="upper right", ncols=2)
save(fig, "02_weight_vs_mpg.png")

# 5.3 MPG by cylinders
fig, ax = plt.subplots(figsize=(7, 4.4))
cyl_vals = by_cyl["mean"]
bars = ax.bar([f"{c}-cyl" for c in cyl_vals.index], cyl_vals.values, width=0.5, color="#2a78d6")
ax.bar_label(bars, labels=[f"{v:.1f}" for v in cyl_vals.values], padding=4, color=INK, fontsize=10,
             fontweight="bold")
ax.set_title("4-cylinder cars average nearly double the MPG of V8s")
ax.title.set_fontsize(11.5)
ax.set_ylabel("Average miles per gallon")
ax.set_ylim(0, cyl_vals.max() * 1.15)
ax.grid(axis="x", visible=False)
save(fig, "03_mpg_by_cylinders.png")

# 5.4 High-efficiency share by era & origin
fig, ax = plt.subplots(figsize=(8, 4.6))
x = np.arange(len(by_era_origin.index))
w = 0.26
for i, o in enumerate(origin_order):
    vals = by_era_origin[o].values * 100
    b = ax.bar(x + (i - 1) * (w + 0.02), vals, width=w, color=COL[o], label=o)
    ax.bar_label(b, labels=[f"{v:.0f}%" for v in vals], padding=3, fontsize=8.5, color=INK2)
ax.set_xticks(x, by_era_origin.index)
ax.set_title("The share of 30+ MPG models climbed in every region")
ax.title.set_fontsize(11.5)
ax.set_ylabel("Share of models with 30+ MPG (%)")
ax.set_xlabel("Model-year era")
ax.set_ylim(0, 90)
ax.grid(axis="x", visible=False)
ax.legend(loc="upper left", ncols=3)
save(fig, "04_high_efficiency_share.png")

# 5.5 Correlation heatmap (diverging: blue = negative, orange = positive, neutral midpoint)
from matplotlib.colors import LinearSegmentedColormap
cmap = LinearSegmentedColormap.from_list("div", ["#2a78d6", "#f1f0ec", "#eb6834"])
labels = ["MPG", "Weight", "Displacement", "Horsepower", "Acceleration"]
fig, ax = plt.subplots(figsize=(6.4, 5.2))
im = ax.imshow(corr.values, cmap=cmap, vmin=-1, vmax=1)
ax.set_xticks(range(5), labels, rotation=30, ha="right")
ax.set_yticks(range(5), labels)
ax.grid(False)
for s in ax.spines.values():
    s.set_visible(False)
for i in range(5):
    for j in range(5):
        v = corr.values[i, j]
        ax.text(j, i, f"{v:.2f}", ha="center", va="center", fontsize=9.5,
                color="white" if abs(v) > 0.65 else INK)
ax.set_title("What moves with fuel efficiency?")
cb = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.03)
cb.outline.set_visible(False)
save(fig, "05_correlation_heatmap.png")

# ----------------------------------------------------------------------------
# 6. EXCEL DASHBOARD (formula-driven, native charts)
# ----------------------------------------------------------------------------
wb = Workbook()
dash = wb.active
dash.title = "Dashboard"
data_ws = wb.create_sheet("Data")

cols = ["name", "make", "mpg", "cylinders", "displacement", "horsepower", "weight_lbs",
        "acceleration", "model_year", "origin", "power_to_weight", "efficiency_class"]
heads = ["Name", "Make", "MPG", "Cylinders", "Displacement", "Horsepower", "Weight (lbs)",
         "Acceleration", "Model Year", "Origin", "HP per 1,000 lbs", "Efficiency Class"]
NAVY, TEAL, LIGHT = "1F3A5F", "1BAF7A", "EEF2F7"
hdr_fill = PatternFill("solid", fgColor=NAVY)
hdr_font = Font(bold=True, color="FFFFFF", name="Calibri")
data_ws.append(heads)
for c in range(1, len(heads) + 1):
    cell = data_ws.cell(1, c)
    cell.fill, cell.font = hdr_fill, hdr_font
for row in df[cols].itertuples(index=False):
    data_ws.append([str(v) if c == "efficiency_class" else (v.item() if hasattr(v, "item") else v)
                    for c, v in zip(cols, row)])
for i, w_ in enumerate([30, 14, 8, 10, 13, 12, 13, 13, 11, 10, 16, 16], 1):
    data_ws.column_dimensions[get_column_letter(i)].width = w_
data_ws.freeze_panes = "A2"
N = len(df) + 1                                    # last data row

rng = lambda col: f"Data!${col}$2:${col}${N}"
MPG, HP, WT, YR, ORG, CYL = rng("C"), rng("F"), rng("G"), rng("I"), rng("J"), rng("D")

dash.sheet_view.showGridLines = False
dash["A1"] = "Automotive Fuel Efficiency & Performance Dashboard"
dash["A1"].font = Font(size=18, bold=True, color=NAVY)
dash["A2"] = "398 vehicles, model years 1970-1982 | all figures recalculate from the Data sheet"
dash["A2"].font = Font(size=10, italic=True, color="666666")

kpis = [("Vehicles analysed", f"=COUNT({MPG})", "0"),
        ("Average MPG", f"=AVERAGE({MPG})", "0.0"),
        ("Avg horsepower", f"=AVERAGE({HP})", "0"),
        ("Avg weight (lbs)", f"=AVERAGE({WT})", "#,##0"),
        ("Models with 30+ MPG", f'=COUNTIF({MPG},">=30")/COUNT({MPG})', "0.0%")]
for i, (lab, f, fmt) in enumerate(kpis):
    col = 1 + i * 2
    dash.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
    dash.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
    a, b = dash.cell(4, col, lab), dash.cell(5, col, f)
    a.font, a.alignment = Font(size=9, color="666666"), Alignment(horizontal="center")
    b.font, b.alignment, b.number_format = Font(size=18, bold=True, color=NAVY), Alignment(horizontal="center"), fmt
    for r in (4, 5):
        for c in (col, col + 1):
            dash.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT)

def table_header(row, labels):
    for j, t in enumerate(labels, 1):
        c = dash.cell(row, j, t)
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, Alignment(horizontal="center")

dash["A7"] = "Performance by origin"
dash["A7"].font = Font(size=12, bold=True, color=NAVY)
table_header(8, ["Origin", "Vehicles", "Avg MPG", "Avg HP", "Avg weight (lbs)", "% 30+ MPG"])
for i, o in enumerate(origin_order):
    r = 9 + i
    dash.cell(r, 1, o)
    dash.cell(r, 2, f'=COUNTIFS({ORG},$A{r})')
    dash.cell(r, 3, f'=AVERAGEIFS({MPG},{ORG},$A{r})').number_format = "0.0"
    dash.cell(r, 4, f'=AVERAGEIFS({HP},{ORG},$A{r})').number_format = "0"
    dash.cell(r, 5, f'=AVERAGEIFS({WT},{ORG},$A{r})').number_format = "#,##0"
    dash.cell(r, 6, f'=COUNTIFS({ORG},$A{r},{MPG},">=30")/B{r}').number_format = "0.0%"

dash["A14"] = "Average MPG by model year"
dash["A14"].font = Font(size=12, bold=True, color=NAVY)
table_header(15, ["Model year"] + origin_order + ["All"])
years = sorted(int(y) for y in df['model_year'].unique())   # 1981 has no records in this dataset
for i, y in enumerate(years):
    r = 16 + i
    dash.cell(r, 1, y)
    for j, o in enumerate(origin_order):
        col = get_column_letter(2 + j)
        dash.cell(r, 2 + j, f'=IFERROR(AVERAGEIFS({MPG},{YR},$A{r},{ORG},{col}$15),NA())').number_format = "0.0"
    dash.cell(r, 5, f'=AVERAGEIFS({MPG},{YR},$A{r})').number_format = "0.0"
last_year_row = 16 + len(years) - 1

dash["A31"] = "Average MPG by cylinders"
dash["A31"].font = Font(size=12, bold=True, color=NAVY)
table_header(32, ["Cylinders", "Vehicles", "Avg MPG"])
for i, c in enumerate([4, 6, 8]):
    r = 33 + i
    dash.cell(r, 1, c)
    dash.cell(r, 2, f'=COUNTIFS({CYL},$A{r})')
    dash.cell(r, 3, f'=AVERAGEIFS({MPG},{CYL},$A{r})').number_format = "0.0"

for i, w_ in enumerate([16, 12, 12, 12, 17, 12, 4, 4, 4, 4], 1):
    dash.column_dimensions[get_column_letter(i)].width = w_

# Native charts
lc = LineChart()
lc.title, lc.height, lc.width = "Average MPG by model year", 8.5, 17
lc.y_axis.title, lc.x_axis.title = "MPG", "Model year"
lc.add_data(Reference(dash, min_col=2, max_col=4, min_row=15, max_row=last_year_row), titles_from_data=True)
lc.set_categories(Reference(dash, min_col=1, min_row=16, max_row=last_year_row))
for s, color in zip(lc.series, ["2A78D6", "EB6834", "1BAF7A"]):
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 22000
    s.smooth = False
lc.x_axis.delete = lc.y_axis.delete = False
dash.add_chart(lc, "H4")

bc = BarChart()
bc.type, bc.title, bc.height, bc.width = "col", "Average MPG by origin", 8.5, 17
bc.add_data(Reference(dash, min_col=3, min_row=8, max_row=11), titles_from_data=True)
bc.set_categories(Reference(dash, min_col=1, min_row=9, max_row=11))
bc.series[0].graphicalProperties.solidFill = "2A78D6"
bc.legend = None
bc.x_axis.delete = bc.y_axis.delete = False
dash.add_chart(bc, "H22")

wb.save(OUT / "Auto_Efficiency_Dashboard.xlsx")

# ----------------------------------------------------------------------------
# 7. REPORT TO CONSOLE
# ----------------------------------------------------------------------------
print(json.dumps(metrics, indent=2, default=str))
print("Done. Charts -> charts/, Excel -> output/Auto_Efficiency_Dashboard.xlsx")
